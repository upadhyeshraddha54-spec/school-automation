import os
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from schoolops_sync.config import load_config  # noqa: E402
from schoolops_sync.db import ALL_TABLES, connect, init_db  # noqa: E402
from schoolops_sync.synthetic import HEADINGS, generate  # noqa: E402

HERE = Path(__file__).parent


@pytest.fixture
def folder(tmp_path):
    f = tmp_path / "school"
    generate(f, per_section=12)
    return f


@pytest.fixture
def cfg(folder):
    c = load_config(HERE / "test_config.toml")
    if os.environ.get("TEST_DATABASE_URL"):
        c.database_url = os.environ["TEST_DATABASE_URL"]
    c.watch_folder = folder
    return c


@pytest.fixture
def conn(cfg):
    c = connect(cfg.database_url)
    init_db(c)
    c.execute("TRUNCATE " + ", ".join(ALL_TABLES) + " RESTART IDENTITY CASCADE")
    yield c
    c.close()


def edit_cell(path: Path, sheet: str, excel_row: int, heading: str, value):
    """Simulate the office editing a SYNTHETIC workbook."""
    wb = load_workbook(path)
    ws = wb[sheet]
    ws.cell(row=excel_row, column=HEADINGS.index(heading) + 1, value=value)
    wb.save(path)


def cell(path: Path, sheet: str, excel_row: int, heading: str):
    wb = load_workbook(path, read_only=True)
    v = list(wb[sheet].iter_rows(min_row=excel_row, max_row=excel_row, values_only=True))[0][HEADINGS.index(heading)]
    wb.close()
    return v


def delete_row(path: Path, sheet: str, excel_row: int):
    wb = load_workbook(path)
    wb[sheet].delete_rows(excel_row)
    wb.save(path)
