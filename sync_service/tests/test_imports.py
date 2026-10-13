import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from conftest import edit_cell
from schoolops_sync.api import create_app

STD1, STD10 = "English - Std 1.xlsx", "Semi-English - Std 10.xlsx"


@pytest.fixture
def client(conn, cfg, tmp_path):
    cfg.upload_folder = str(tmp_path / "uploads")
    with TestClient(create_app(cfg, start_watcher=False)) as c:
        yield c


def files(folder, *names):
    return [("files", (n, (folder / n).read_bytes(), "application/vnd.ms-excel")) for n in names]


def n_students(conn):
    return conn.execute("SELECT count(*) AS n FROM students").fetchone()["n"]


def test_preview_validates_and_writes_nothing(client, conn, folder):
    r = client.post("/imports/preview", files=files(folder, STD1, STD10)).json()
    assert r["token"] and len(r["files"]) == 2
    f1 = r["files"][0]
    assert f1["importable"] and f1["students_in_file"] == 24
    assert all(c["ok"] for c in f1["columns"])
    assert f1["plan"]["new_students"] == 24 and len(f1["rows"]) == 8
    assert [s["sheet"] for s in f1["sheets"]] == ["1 Peacock", "1 Swan"]
    assert n_students(conn) == 0  # preview only


def test_import_happens_only_after_confirm(client, conn, folder, cfg):
    r = client.post("/imports/preview", files=files(folder, STD1, STD10)).json()
    done = client.post("/imports/confirm", json={"token": r["token"], "files": [STD1]}).json()
    assert done["files"][0]["status"] == "applied"
    assert n_students(conn) == 24  # only the chosen file
    assert client.post("/imports/confirm", json={"token": r["token"], "files": [STD10]}).status_code == 400  # used up


def test_reimport_keeps_calls_and_updates_fees(client, conn, folder):
    r = client.post("/imports/preview", files=files(folder, STD1)).json()
    client.post("/imports/confirm", json={"token": r["token"], "files": [STD1]})
    sid = conn.execute("SELECT s.id FROM students s JOIN current_fees f ON f.student_id=s.id "
                       "WHERE s.source_sheet='1 Peacock' AND s.sr_no='1'").fetchone()["id"]
    client.post("/calls", json={"student_id": sid, "outcome": "will_pay"})
    edit_cell(folder / STD1, "1 Peacock", 3, "Balance", 0)
    r2 = client.post("/imports/preview", files=files(folder, STD1)).json()
    assert r2["files"][0]["plan"]["new_students"] == 0 and r2["files"][0]["plan"]["fee_updates"] == 1
    client.post("/imports/confirm", json={"token": r2["token"], "files": [STD1]})
    assert n_students(conn) == 24
    assert client.get("/calls", params={"student_id": sid}).json()["total"] == 1


def test_old_xls_and_real_files_are_blocked_in_test_mode(client, conn, tmp_path):
    real = Workbook()
    ws = real.active
    ws.title = "1 Peacock"
    ws.append(["SHREE REAL SCHOOL"])
    ws.append(["Sr. No.", "Student Name", "Class", "Contact", "Total Amt", "Disc. Amt", "Paid Amt", "Balance"])
    buf = io.BytesIO()
    real.save(buf)
    r = client.post("/imports/preview", files=[
        ("files", ("Std 1.xls", b"old excel", "application/vnd.ms-excel")),
        ("files", ("English - Std 1.xlsx", buf.getvalue(), "application/vnd.ms-excel")),
    ]).json()
    assert r["token"] is None
    assert ".xlsx" in r["files"][0]["problem"]
    assert "synthetic" in r["files"][1]["problem"].lower()
    assert n_students(conn) == 0


def test_missing_column_is_reported(client, folder):
    p = folder / STD1
    wb = load_workbook(p)
    wb["1 Swan"]["H2"] = "Remaining"
    wb.save(p)
    f = client.post("/imports/preview", files=files(folder, STD1)).json()["files"][0]
    bal = next(c for c in f["columns"] if c["field"] == "balance")
    assert bal["ok"] is False and bal["missing_in"] == ["1 Swan"]
    assert f["sheets_failed"] == ["1 Swan"] and f["importable"]  # 1 Peacock is still fine


def test_bad_token_and_file_name_are_rejected(client, folder):
    assert client.post("/imports/confirm", json={"token": "../../etc", "files": [STD1]}).status_code == 400
    r = client.post("/imports/preview", files=files(folder, STD1)).json()
    assert client.post("/imports/confirm", json={"token": r["token"], "files": ["../config.toml"]}).status_code == 400
    assert client.delete(f"/imports/{r['token']}").status_code == 200


def test_unconfirmed_mapping_blocks_real_imports(client, cfg, folder):
    cfg.test_mode = False
    cfg.mapping_confirmed = False
    r = client.post("/imports/preview", files=files(folder, STD1)).json()
    assert r["token"] is None and "mapping" in r["files"][0]["problem"]
