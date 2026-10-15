"""OneDrive via Microsoft Graph, tested against a fake Graph server (no network, synthetic files)."""
import io
import os
from pathlib import Path
from urllib.parse import parse_qs

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from schoolops_sync.api import create_app
from schoolops_sync.config import SYNTHETIC_MARKER, NotAllowed
from schoolops_sync.onedrive import GraphClient, NotConnected, OneDriveConfig, sync_from_onedrive
from schoolops_sync.synthetic import HEADINGS, generate


class FakeGraph:
    """Minimal Graph: folder listing with paging, download URLs, throttling. Records every request."""

    def __init__(self, page_size=2):
        self.items, self.requests, self.page_size = {}, [], page_size
        self.throttle_next, self.unauthorized = 0, False

    def add_file(self, iid, name, data, parent="root"):
        self.items[iid] = {"id": iid, "name": name, "parent": parent, "data": data, "ctag": f"c-{iid}-1", "v": 1}

    def add_folder(self, iid, name, parent="root"):
        self.items[iid] = {"id": iid, "name": name, "parent": parent, "folder": True}

    def change(self, iid, data):
        it = self.items[iid]
        it["v"] += 1
        it["data"], it["ctag"] = data, f"c-{iid}-{it['v']}"

    def _json(self, it):
        base = {"id": it["id"], "name": it["name"], "parentReference": {"driveId": "D1"}}
        if it.get("folder"):
            return {**base, "folder": {"childCount": 1}}
        return {**base, "file": {}, "cTag": it["ctag"], "size": len(it["data"]),
                "lastModifiedDateTime": "2026-10-15T05:00:00Z",
                "@microsoft.graph.downloadUrl": f"https://download.test/{it['id']}"}

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.requests.append((request.method, str(request.url)))
        assert request.method == "GET", "the app must never write to OneDrive"
        if self.throttle_next:
            self.throttle_next -= 1
            return httpx.Response(429, headers={"Retry-After": "0"})
        if request.url.host == "download.test":
            assert "authorization" not in request.headers  # pre-authenticated link: no token sent
            return httpx.Response(200, content=self.items[request.url.path.strip("/")]["data"])
        if self.unauthorized:
            return httpx.Response(401)
        assert request.headers["authorization"] == "Bearer TEST"
        path = request.url.path
        if path.endswith("/me"):
            return httpx.Response(200, json={"displayName": "Test Office", "userPrincipalName": "office@test.example"})
        parent = "root" if path.endswith("root:/School/Fees:/children") else path.split("/items/")[1].split("/")[0]
        kids = [self._json(i) for i in self.items.values() if i["parent"] == parent]
        skip = int(parse_qs(request.url.query.decode()).get("skip", ["0"])[0])
        page = kids[skip: skip + self.page_size]
        body = {"value": page}
        if skip + self.page_size < len(kids):
            body["@odata.nextLink"] = str(request.url.copy_with(params={"skip": skip + self.page_size}))
        return httpx.Response(200, json=body)


@pytest.fixture
def setup(cfg, tmp_path):
    src = tmp_path / "source"
    generate(src)
    g = FakeGraph()
    g.add_folder("F1", "Semi-English")
    for i, p in enumerate(sorted(src.glob("*.xlsx"))):
        g.add_file(f"I{i}", p.name, p.read_bytes(), parent="F1" if p.name.startswith("Semi") else "root")
    mirror = tmp_path / "mirror"
    mirror.mkdir()
    (mirror / SYNTHETIC_MARKER).write_text("test")
    cfg.watch_folder = mirror
    oc = OneDriveConfig.from_raw({"enabled": True, "folder_path": "School/Fees", "client_id": "x"})
    client = lambda: GraphClient(lambda: "TEST", transport=httpx.MockTransport(g.handler), sleep=lambda s: None)  # noqa: E731
    return g, oc, client, mirror, src


def count(conn, t, where="TRUE"):
    return conn.execute(f"SELECT count(*) AS n FROM {t} WHERE {where}").fetchone()["n"]


def edited(data: bytes, sheet, row, heading, value) -> bytes:
    wb = load_workbook(io.BytesIO(data))
    wb[sheet].cell(row=row, column=HEADINGS.index(heading) + 1, value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_initial_download_and_sync_then_nothing_to_do(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    res = sync_from_onedrive(conn, cfg, oc, client())
    assert sorted(res["downloaded_or_renamed"]) == sorted(
        ["English - Nursery.xlsx", "English - Std 1.xlsx", "English - Std 10.xlsx",
         "Semi-English/Semi-English - Std 1.xlsx", "Semi-English/Semi-English - Std 10.xlsx"])
    assert all(s["status"] == "applied" for s in res["sync"])
    assert count(conn, "students") == 84
    assert all(m == "GET" for m, _ in g.requests)
    assert (mirror / "English - Std 1.xlsx").stat().st_mode & 0o222 == 0  # mirror copies are read-only
    again = sync_from_onedrive(conn, cfg, oc, client())
    assert again["downloaded_or_renamed"] == [] and again["unchanged"] == 5 and again["sync"] == []
    assert not any("download.test" in u for _, u in g.requests[len(g.requests) - again["graph_requests"]:])


def test_changed_file_is_downloaded_and_payment_counted_once(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    sync_from_onedrive(conn, cfg, oc, client())
    it = next(i for i in g.items.values() if i["name"] == "English - Std 1.xlsx")
    s = conn.execute("""SELECT s.id, f.balance, f.paid FROM students s JOIN current_fees f ON f.student_id=s.id
                        WHERE s.source_sheet='1 Peacock' AND s.sr_no='1'""").fetchone()
    data = edited(it["data"], "1 Peacock", 3, "Paid Amt", float(s["paid"] + 2000))
    g.change(it["id"], edited(data, "1 Peacock", 3, "Balance", float(s["balance"] - 2000)))
    res = sync_from_onedrive(conn, cfg, oc, client())
    assert res["downloaded_or_renamed"] == ["English - Std 1.xlsx"] and res["sync"][0]["fee_updates"] == 1
    assert count(conn, "payments") == 1
    assert conn.execute("SELECT amount FROM payments").fetchone()["amount"] == 2000
    sync_from_onedrive(conn, cfg, oc, client())
    assert count(conn, "payments") == 1


def test_rename_move_and_removal(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    sync_from_onedrive(conn, cfg, oc, client())
    iid = next(k for k, i in g.items.items() if i.get("name") == "English - Std 10.xlsx")
    g.items[iid]["name"] = "English - Std 10 (2026-27).xlsx"
    n_req = len(g.requests)
    res = sync_from_onedrive(conn, cfg, oc, client())
    assert res["renamed"] == [{"from": "English - Std 10.xlsx", "to": "English - Std 10 (2026-27).xlsx"}]
    assert not any("download.test" in u for _, u in g.requests[n_req:])  # no re-download for a rename
    assert (mirror / "English - Std 10 (2026-27).xlsx").exists() and not (mirror / "English - Std 10.xlsx").exists()
    assert count(conn, "students") == 84 and res["sync"][0]["new_students"] == 0
    # removed in OneDrive -> local copy moved aside, students kept
    del g.items[iid]
    res = sync_from_onedrive(conn, cfg, oc, client())
    assert res["removed"] == ["English - Std 10 (2026-27).xlsx"]
    assert list((mirror / ".removed").iterdir()) and count(conn, "students") == 84


def test_throttling_and_paging(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    g.throttle_next = 2
    res = sync_from_onedrive(conn, cfg, oc, client())
    assert len(res["downloaded_or_renamed"]) == 5  # 3 root files over 2 pages + subfolder


def test_real_folder_is_refused_before_any_request(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    cfg.test_mode = False
    cfg.mapping_confirmed = False
    with pytest.raises(NotAllowed):
        sync_from_onedrive(conn, cfg, oc, client())
    assert g.requests == []


def test_non_synthetic_file_is_blocked_in_test_mode(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    wb = Workbook()
    ws = wb.active
    ws.title = "1 Peacock"
    ws.append(["REAL SCHOOL"])
    ws.append(HEADINGS)
    buf = io.BytesIO()
    wb.save(buf)
    g.add_file("REAL", "English - Std 2.xlsx", buf.getvalue())
    res = sync_from_onedrive(conn, cfg, oc, client())
    assert res["blocked_not_synthetic"] == ["English - Std 2.xlsx"]
    assert not (mirror / "English - Std 2.xlsx").exists()


def test_expired_sign_in_is_reported(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    g.unauthorized = True
    with pytest.raises(NotConnected):
        sync_from_onedrive(conn, cfg, oc, client())
    st = conn.execute("SELECT value FROM settings WHERE key='onedrive_status'").fetchone()["value"]
    assert st["ok"] is False and "sign" in st["error"].lower()


def test_api_status_and_sync_now(conn, cfg, setup):
    g, oc, client, mirror, _ = setup
    cfg.onedrive = {"enabled": True, "folder_path": "School/Fees", "client_id": "x"}
    with TestClient(create_app(cfg, start_watcher=False, graph_client_factory=client)) as c:
        r = c.post("/onedrive/sync")
        assert r.status_code == 200 and len(r.json()["sync"]) == 5
        st = c.get("/onedrive/status").json()
        assert st["enabled"] and st["files_tracked"] == 5 and st["last"]["ok"] is True
    assert count(conn, "students") == 84


def test_folder_url_variants():
    from schoolops_sync.onedrive import folder_children_url
    assert folder_children_url(OneDriveConfig(folder_path="School/Fees 2026")).endswith("/me/drive/root:/School/Fees%202026:/children")
    assert "/drives/D9/items/X1/children" in folder_children_url(OneDriveConfig(drive_id="D9", folder_item_id="X1"))
    with pytest.raises(ValueError):
        OneDriveConfig.from_raw({"folder_item_id": "X"})


def test_verify_reports_account_and_headings_only(conn, cfg, setup):
    from schoolops_sync.onedrive import status, verify
    g, oc, client, mirror, _ = setup
    assert status(conn, cfg, oc)["verified"] is None  # never "live" before a real check
    res = verify(conn, cfg, oc, client(), sample=2)
    assert res["account"] == "office@test.example" and res["excel_files"] == 5 and res["sampled"] == 2
    assert "student_name" in res["inspection"]["mapping"]
    assert count(conn, "students") == 0  # verifying never syncs
    assert list(mirror.glob("*.xlsx")) == []
    assert all(m == "GET" for m, _ in g.requests)
    assert status(conn, cfg, oc)["verified"]["account"] == "office@test.example"
