from decimal import Decimal

import pytest
from pathlib import Path

from openpyxl import Workbook

from conftest import edit_cell
from schoolops_sync.reader import normalize_medium, normalize_std, read_workbook, sha256_file


def test_reads_class_sheets_and_skips_title_and_summary(cfg, folder):
    pf = read_workbook(folder / "English - Std 1.xlsx", cfg)
    assert not pf.fatal
    assert pf.sheets_read == ["1 Peacock", "1 Swan"]
    assert pf.sheets_skipped == ["Summary"]
    assert len(pf.rows) == 24
    r = pf.rows[0]
    assert (r.medium, r.std, r.section, r.sr_no) == ("English", "1", "Peacock", "1")
    assert r.balance == r.total_fee - r.discount - r.paid
    assert all(x.valid for x in pf.rows)


def test_semi_english_and_nursery_classified(cfg, folder):
    semi = read_workbook(folder / "Semi-English - Std 10.xlsx", cfg)
    assert {(r.medium, r.std, r.section) for r in semi.rows} == {("Semi-English", "10", "Lotus")}
    nur = read_workbook(folder / "English - Nursery.xlsx", cfg)
    assert {r.std for r in nur.rows} == {"Nursery"}


def test_broken_xls_fails_loudly(cfg, folder):
    p = folder / "Std 1.xls"
    p.write_bytes(b"not really excel")
    pf = read_workbook(p, cfg)
    assert pf.fatal and pf.issues[0].code == "unreadable_file" and ".xlsx" in pf.issues[0].message


def test_real_xls_is_read_through_a_converted_copy(cfg, folder, tmp_path):
    import shutil
    import subprocess
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed; cannot create an .xls test file")
    src = folder / "English - Std 1.xlsx"
    out = tmp_path / "xls"
    subprocess.run(["soffice", "--headless", "--convert-to", "xls", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=120)
    xls = folder / "English - Std 1 old.xls"
    shutil.move(out / "English - Std 1.xls", xls)
    before = sha256_file(xls)
    pf = read_workbook(xls, cfg)
    assert not pf.fatal and pf.sheets_read == ["1 Peacock", "1 Swan"] and len(pf.rows) == 24
    assert pf.rows[0].contact == "9000001000" and pf.rows[0].balance is not None
    assert any(i.code == "converted_xls" for i in pf.issues)
    assert sha256_file(xls) == before  # original untouched


def test_changed_layout_fails_the_sheet_not_the_whole_file(cfg, folder):
    p = folder / "English - Std 1.xlsx"
    from openpyxl import load_workbook
    wb = load_workbook(p)
    wb["1 Swan"]["H2"] = "Remaining"  # Balance heading renamed
    wb.save(p)
    pf = read_workbook(p, cfg)
    assert pf.sheets_failed == ["1 Swan"] and pf.sheets_read == ["1 Peacock"]
    assert any(i.code == "missing_headings" for i in pf.issues)
    assert {r.sheet for r in pf.rows} == {"1 Peacock"}


def test_amount_text_negative_and_broken_phone(cfg, folder):
    p = folder / "English - Std 1.xlsx"
    edit_cell(p, "1 Peacock", 3, "Balance", "₹12,500")
    edit_cell(p, "1 Peacock", 4, "Balance", -100)
    edit_cell(p, "1 Peacock", 5, "Contact", "8E+09")
    pf = read_workbook(p, cfg)
    rows = {r.row_number: r for r in pf.rows if r.sheet == "1 Peacock"}
    assert rows[3].balance == Decimal("12500.00") and rows[3].valid
    assert not rows[4].valid and any(i.code == "bad_amount" and i.row == 4 for i in pf.issues)
    assert rows[5].contact is None and any(i.code == "phone" and i.row == 5 for i in pf.issues)


def test_duplicate_rows_are_not_synced(cfg, folder):
    p = folder / "English - Std 1.xlsx"
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb["1 Peacock"]
    ws.append([v.value for v in ws[3]])  # exact copy of the first student, same Sr. No.
    wb.save(p)
    pf = read_workbook(p, cfg)
    dup = [r for r in pf.rows if r.name == pf.rows[0].name and r.sheet == "1 Peacock"]
    assert len(dup) == 2 and not any(r.valid for r in dup)
    assert any(i.code == "duplicate_student" for i in pf.issues)


def test_reading_never_changes_the_original_file(cfg, folder):
    p = folder / "English - Std 10.xlsx"
    before = (sha256_file(p), p.stat().st_mtime_ns)
    read_workbook(p, cfg)
    assert (sha256_file(p), p.stat().st_mtime_ns) == before


def test_formula_without_saved_value_is_flagged(cfg, folder):
    p = folder / "English - Std 1.xlsx"
    edit_cell(p, "1 Peacock", 3, "Balance", "=E3-F3-G3")  # openpyxl saves no cached result
    pf = read_workbook(p, cfg)
    assert any(i.code == "missing_balance" and i.row == 3 for i in pf.issues)


def test_std_and_medium_normalisation():
    assert [normalize_std(x) for x in ["I", "x", "Std 7", "10th", "Nursery", "Jr KG", "UKG", "11"]] == \
           ["1", "10", "7", "10", "Nursery", "LKG", "UKG", None]
    assert normalize_medium("Semi English") == "Semi-English"
    assert normalize_medium("ENGLISH") == "English"
    assert normalize_medium("Marathi") is None


def test_empty_workbook_is_fatal(cfg, folder):
    p = folder / "Empty.xlsx"
    Workbook().save(p)
    assert read_workbook(p, cfg).fatal
