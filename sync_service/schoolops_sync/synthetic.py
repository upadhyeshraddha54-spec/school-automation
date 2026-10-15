"""Generate a folder of FAKE school fee workbooks for testing.

Names are random combinations, phone numbers start with 90000 (not real), and the
folder gets a marker file so test mode can only ever run on synthetic data.
"""
from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook

from .config import SYNTHETIC_MARKER

FIRST = ["Aarav", "Isha", "Rohan", "Sanika", "Vedant", "Tanvi", "Om", "Anaya", "Kabir", "Mitali",
         "Pranav", "Riya", "Arjun", "Sakshi", "Yash", "Neha", "Siddharth", "Pooja", "Aditya", "Gauri"]
LAST = ["Kulkarni", "Deshmukh", "Jadhav", "More", "Pawar", "Shinde", "Gaikwad", "Bhosale", "Joshi",
        "Patil", "Chavan", "Salunkhe", "Wagh", "Thorat", "Mane", "Lokhande", "Kale", "Ghule"]
HEADINGS = ["Sr. No.", "Student Name", "Class", "Contact", "Total Amt", "Disc. Amt", "Paid Amt", "Balance"]

# medium -> {file label: [sections]}
LAYOUT = {
    "English": {"Nursery": ["Rose"], "Std 1": ["Peacock", "Swan"], "Std 10": ["Peacock", "Swan"]},
    "Semi-English": {"Std 1": ["Lotus"], "Std 10": ["Lotus"]},
}


def make_sheet(ws, std_label: str, section: str, students: list[dict], school="SYNTHETIC TEST SCHOOL"):
    ws.title = f"{std_label} {section}"
    ws["A1"] = school  # title row, like the real files
    ws.append(HEADINGS)
    for i, s in enumerate(students, 1):
        ws.append([i, s["name"], f"{std_label}-{section}", s["contact"], s["total"], s["disc"], s["paid"],
                   s["total"] - s["disc"] - s["paid"]])


def fake_students(rng: random.Random, n: int, start_phone: int) -> list[dict]:
    names, out = set(), []
    while len(out) < n:
        name = f"{rng.choice(LAST)} {rng.choice(FIRST)}"
        if name in names:
            continue
        names.add(name)
        total = rng.choice([18000, 25000, 30000])
        disc = rng.choice([0, 0, 0, 1000])
        paid = rng.choice([0, 5000, 7500, 10000, total - disc])
        out.append({"name": name, "contact": str(start_phone + len(out)), "total": total, "disc": disc, "paid": paid})
    return out


# Synthetic siblings: (medium, file label, section, Sr. No.) pairs that share a parent phone.
SIBLINGS = [
    (("English", "Std 1", "Swan", 8), ("English", "Std 10", "Peacock", 2)),
    (("English", "Std 1", "Swan", 9), ("English", "Std 10", "Peacock", 3)),
    (("English", "Std 1", "Swan", 10), ("English", "Std 10", "Peacock", 4)),
    (("Semi-English", "Std 1", "Lotus", 2), ("English", "Std 10", "Swan", 6)),
]


def generate(folder: str | Path, per_section: int = 12, seed: int = 7, siblings: bool = True) -> list[Path]:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / SYNTHETIC_MARKER).write_text("Synthetic test data only. Not real students.\n")
    rng = random.Random(seed)
    phone = 9000000000
    data: dict[tuple, list[dict]] = {}
    for medium, stds in LAYOUT.items():
        for label, sections in stds.items():
            for sec in sections:
                data[(medium, label, sec)] = fake_students(rng, per_section, phone)
                phone += 1000
    if siblings and per_section >= 10:
        for (m1, l1, s1, i1), (m2, l2, s2, i2) in SIBLINGS:
            a, b = data[(m1, l1, s1)][i1 - 1], data[(m2, l2, s2)][i2 - 1]
            b["contact"] = a["contact"]                      # same parent phone
            b["name"] = a["name"].split()[0] + " " + b["name"].split()[1]  # same surname
    files = []
    for medium, stds in LAYOUT.items():
        for label, sections in stds.items():
            wb = Workbook()
            wb.remove(wb.active)
            std_label = label.replace("Std ", "")
            for sec in sections:
                make_sheet(wb.create_sheet(), std_label, sec, data[(medium, label, sec)])
            wb.create_sheet("Summary")["A1"] = "Class summary (not a student sheet)"
            p = folder / f"{medium} - {label}.xlsx"
            wb.save(p)
            files.append(p)
    return files
