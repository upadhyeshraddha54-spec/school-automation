"""Read and validate a school fees workbook.

Safety: the original file is never opened by openpyxl. It is copied to a private temp
folder first, and only the copy is read (read_only=True, data_only=True). Nothing is
ever written back to Excel.
"""
from __future__ import annotations

import hashlib
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import Classifier, Config

NUMERIC_FIELDS = ("total_fee", "discount", "paid", "balance")
STD_ORDER = ["Nursery", "LKG", "UKG"] + [str(i) for i in range(1, 11)]
_ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}
_STD_WORDS = {
    "nursery": "Nursery", "nur": "Nursery", "playgroup": "Nursery",
    "lkg": "LKG", "jrkg": "LKG", "juniorkg": "LKG",
    "ukg": "UKG", "srkg": "UKG", "seniorkg": "UKG",
}


@dataclass
class Issue:
    level: str  # "error" | "warning" | "info"
    code: str
    message: str
    sheet: str | None = None
    row: int | None = None


@dataclass
class StudentRow:
    sheet: str
    row_number: int
    name: str
    medium: str | None
    std: str | None
    section: str | None
    sr_no: str | None = None
    student_ext_id: str | None = None
    parent_name: str | None = None
    family_ext_id: str | None = None
    contact: str | None = None
    total_fee: Decimal | None = None
    discount: Decimal | None = None
    paid: Decimal | None = None
    balance: Decimal | None = None
    raw: dict[str, Any] = field(default_factory=dict)
    valid: bool = True  # False -> seen in file but values not trusted (not applied)


@dataclass
class ParsedFile:
    path: Path
    file_hash: str
    rows: list[StudentRow] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    sheets_read: list[str] = field(default_factory=list)
    sheets_skipped: list[str] = field(default_factory=list)
    sheets_failed: list[str] = field(default_factory=list)
    fatal: bool = False

    def add(self, level, code, message, sheet=None, row=None):
        self.issues.append(Issue(level, code, message, sheet, row))


def norm_text(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def norm_key(v: Any) -> str:
    return norm_text(v).casefold()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def normalize_std(value: str) -> str | None:
    v = norm_text(value)
    if not v:
        return None
    compact = re.sub(r"[\s.\-_]", "", v).casefold()
    if compact in _STD_WORDS:
        return _STD_WORDS[compact]
    m = re.fullmatch(r"(?:std|class)?(\d{1,2})(?:st|nd|rd|th)?", compact)
    if m and 1 <= int(m.group(1)) <= 10:
        return str(int(m.group(1)))
    up = re.sub(r"^(STD|CLASS)", "", compact.upper())
    if up in _ROMAN:
        return str(_ROMAN[up])
    return None


def normalize_medium(value: str) -> str | None:
    c = re.sub(r"[\s\-_]", "", norm_text(value)).casefold()
    if c in ("english", "eng", "em"):
        return "English"
    if c in ("semienglish", "semi", "semieng", "sem"):
        return "Semi-English"
    return None


def _number(v: Any) -> tuple[Decimal | None, str | None]:
    """Returns (value, error). Blank -> (None, None)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    if isinstance(v, bool):
        return None, f"unexpected value {v!r}"
    if isinstance(v, (int, float)):
        d = Decimal(str(v))
    else:
        s = re.sub(r"[₹,\s]|rs\.?", "", str(v), flags=re.I)
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None, f"'{v}' is not a number"
    if d < 0:
        return None, f"negative amount {d}"
    return d.quantize(Decimal("0.01")), None


def _phone(v: Any) -> tuple[str | None, str | None]:
    """Returns (phone, warning)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, "no phone number"
    if isinstance(v, float):
        if not v.is_integer():
            return None, f"phone looks wrong: {v}"
        v = int(v)
    s = str(v).strip()
    if re.search(r"\d[eE][+-]?\d", s):
        return None, f"phone is broken in Excel ('{s}'). Format the Contact column as text and re-type it"
    digits = re.sub(r"\D", "", s)
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    if len(digits) != 10:
        return digits or None, f"phone '{s}' is not a 10-digit number"
    return digits, None


def _id_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = norm_text(v)
    return s or None


def _classify(c: Classifier | None, file_stem: str, sheet: str, class_cell: Any) -> str | None:
    if c is None:
        return None
    if c.source == "value":
        return c.regex
    src = {"file": file_stem, "sheet": sheet, "class_column": norm_text(class_cell)}[c.source]
    m = re.search(c.regex, src, flags=re.I)
    if not m:
        return None
    return norm_text(m.group(1) if m.groups() else m.group(0))


def _xls_to_xlsx(src: Path, dst: Path) -> None:
    """Old .xls -> a temporary .xlsx copy (values only). The original is only read."""
    import xlrd  # optional dependency, only needed for .xls
    from openpyxl import Workbook

    book = xlrd.open_workbook(str(src), on_demand=True)
    wb = Workbook()
    wb.remove(wb.active)
    try:
        for sh in book.sheets():
            ws = wb.create_sheet(sh.name[:31])
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        v = None
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        v = None
                    elif cell.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                        v = int(v)
                    row.append(v)
                ws.append(row)
    finally:
        book.release_resources()
    wb.save(dst)


def read_workbook(path: str | Path, cfg: Config) -> ParsedFile:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in (".xlsx", ".xlsm", ".xls"):
        pf = ParsedFile(path=path, file_hash="")
        pf.fatal = True
        pf.add("error", "unsupported_format", f"{path.name}: not an Excel file (.xlsx or .xls).")
        return pf

    with tempfile.TemporaryDirectory(prefix="schoolops-") as tmp:
        copy = Path(tmp) / ("copy" + suffix)
        try:
            shutil.copy2(path, copy)  # read-only copy; the original is never opened for writing
        except OSError as exc:  # e.g. OneDrive file not downloaded yet, or locked
            pf = ParsedFile(path=path, file_hash="")
            pf.fatal = True
            pf.add("error", "file_not_available",
                   f"{path.name}: could not be read right now ({exc.__class__.__name__}). "
                   "If it is in OneDrive, it may still be downloading. It will be tried again.")
            return pf
        pf = ParsedFile(path=path, file_hash=sha256_file(copy))
        if suffix == ".xls":
            try:
                converted = Path(tmp) / "converted.xlsx"
                _xls_to_xlsx(copy, converted)
                copy = converted
                pf.add("info", "converted_xls", f"{path.name}: old .xls format, read through a converted copy.")
            except ImportError:
                pf.fatal = True
                pf.add("error", "xls_reader_missing",
                       f"{path.name}: .xls files need the 'xlrd' package (pip install xlrd), or Save As .xlsx.")
                return pf
            except Exception as exc:
                pf.fatal = True
                pf.add("error", "unreadable_file",
                       f"{path.name}: could not read this .xls file ({exc.__class__.__name__}). Open it and Save As .xlsx.")
                return pf
        try:
            wb = load_workbook(copy, read_only=True, data_only=True)
        except Exception as exc:  # corrupt / half-saved / password protected
            pf.fatal = True
            pf.add("error", "unreadable_file", f"{path.name}: could not open workbook ({exc.__class__.__name__}).")
            return pf
        try:
            for ws in wb.worksheets:
                _read_sheet(pf, cfg, path.stem, ws.title, list(ws.iter_rows(values_only=True)))
        finally:
            wb.close()

    if not pf.sheets_read and not pf.sheets_failed:
        pf.fatal = True
        pf.add("error", "no_class_sheets", f"{path.name}: no sheet has the expected headings.")
    _flag_duplicates(pf)
    return pf


def _read_sheet(pf: ParsedFile, cfg: Config, stem: str, sheet: str, rows: list[tuple]) -> None:
    hidx = cfg.header_row - 1
    # Safety net: if the headings are not on the configured row (e.g. someone inserted a row
    # above them), look in the first 8 rows for the row that has BOTH name and balance headings.
    def _has_core(row) -> bool:
        keys = {norm_key(h) for h in (row or ()) if h is not None}
        return all(any(norm_key(h) in keys for h in cfg.headings(f)) for f in ("student_name", "balance"))
    if not (len(rows) > hidx and _has_core(rows[hidx])):
        for r in range(min(len(rows), 8)):
            if _has_core(rows[r]):
                pf.add("warning", "headings_moved",
                       f"Headings found on row {r + 1} instead of row {cfg.header_row}. Read from there.", sheet)
                hidx = r
                break
    data_start = hidx + (cfg.first_data_row - cfg.header_row)
    header = rows[hidx] if len(rows) > hidx else ()
    found: dict[str, int] = {}
    hnorm = [norm_key(h) for h in header]
    for fld, heading in cfg.columns.items():
        accepted = {norm_key(x) for x in cfg.headings(fld)}
        hits = [i for i, h in enumerate(hnorm) if h and h in accepted]
        if len(hits) > 1:
            pf.add("error", "duplicate_heading", f"Heading '{heading}' (or an accepted variant) appears {len(hits)} times.", sheet)
            pf.sheets_failed.append(sheet)
            return
        if hits:
            found[fld] = hits[0]

    if "student_name" not in found and "balance" not in found:
        pf.sheets_skipped.append(sheet)
        pf.add("info", "sheet_skipped", "Not a class sheet (expected headings not in the heading row).", sheet)
        return
    missing = [cfg.columns[f] for f in cfg.columns if f not in found]
    if missing:
        pf.sheets_failed.append(sheet)
        pf.add("error", "missing_headings",
               f"Missing heading(s) {missing} in row {cfg.header_row}. The file layout changed; nothing from this sheet was synced.",
               sheet)
        return

    pf.sheets_read.append(sheet)
    for r in range(data_start, len(rows)):
        row = rows[r]
        excel_row = r + 1
        if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        get = lambda f: row[found[f]] if f in found and found[f] < len(row) else None  # noqa: E731
        name = norm_text(get("student_name"))
        if not name:
            pf.add("warning", "row_without_name", "Row has values but no student name; ignored.", sheet, excel_row)
            continue
        if re.match(r"^(grand\s*)?total\b", name, re.I):
            continue

        cell_class = get("class_label")
        std_raw = _classify(cfg.std, stem, sheet, cell_class)
        sec = _classify(cfg.section, stem, sheet, cell_class)
        med_raw = _classify(cfg.medium, stem, sheet, cell_class)
        std = normalize_std(std_raw) if std_raw else None
        medium = normalize_medium(med_raw) if med_raw else None

        sr_out = _id_text(get("sr_no"))
        s = StudentRow(sheet=sheet, row_number=excel_row, name=name, medium=medium, std=std,
                       section=sec.title() if sec else None, sr_no=sr_out,
                       student_ext_id=_id_text(get("student_id")),
                       parent_name=norm_text(get("parent_name")) or None,
                       family_ext_id=_id_text(get("family_id")),
                       raw={cfg.columns[f]: _jsonable(get(f)) for f in found})

        for label, val, fixed in (("class", std_raw, std), ("medium", med_raw, medium)):
            if fixed is None:
                s.valid = False
                pf.add("error", f"unknown_{label}",
                       f"Could not work out the {label} for '{name}' (got '{val or ''}').", sheet, excel_row)
        if cfg.section and not sec:
            s.valid = False
            pf.add("error", "unknown_section", f"Could not work out the section for '{name}'.", sheet, excel_row)

        for f in NUMERIC_FIELDS:
            if f in found:
                val, err = _number(get(f))
                if err:
                    s.valid = False
                    pf.add("error", "bad_amount", f"{cfg.columns[f]}: {err}.", sheet, excel_row)
                setattr(s, f, val)
        if s.balance is None and s.valid:
            s.valid = False
            pf.add("error", "missing_balance",
                   f"'{name}' has no Balance value (blank or a formula without a saved result).", sheet, excel_row)

        if "contact" in found:
            s.contact, warn = _phone(get("contact"))
            if warn:
                pf.add("warning", "phone", f"'{name}': {warn}.", sheet, excel_row)

        if s.valid and None not in (s.total_fee, s.paid, s.balance):
            expected = s.total_fee - (s.discount or 0) - s.paid
            if expected != s.balance:
                pf.add("info", "balance_differs",
                       f"'{name}': Total − Discount − Paid = {expected}, but Balance says {s.balance}. "
                       "Synced as written (may include installments or old dues).", sheet, excel_row)
        pf.rows.append(s)


def _jsonable(v: Any):
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def identity_key(row: StudentRow) -> str:
    """Never guessed: a real ID column if the school has one, otherwise the exact
    combination of medium + class + section + Sr. No. + name as written in Excel."""
    if row.student_ext_id:
        return f"id:{norm_key(row.student_ext_id)}"
    return "|".join(norm_key(x) for x in (row.medium, row.std, row.section, row.sr_no, row.name))


def _flag_duplicates(pf: ParsedFile) -> None:
    seen: dict[str, list[StudentRow]] = {}
    for r in pf.rows:
        seen.setdefault(identity_key(r), []).append(r)
    for rows in seen.values():
        if len(rows) > 1:
            where = ", ".join(f"{r.sheet} row {r.row_number}" for r in rows)
            for r in rows:
                r.valid = False
            pf.add("error", "duplicate_student",
                   f"'{rows[0].name}' appears {len(rows)} times ({where}). None of these rows were synced.",
                   rows[0].sheet, rows[0].row_number)
