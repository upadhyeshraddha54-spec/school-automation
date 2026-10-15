"""Look at the school's real workbooks ON THIS COMPUTER and propose the column mapping.

Only headings, sheet names and counts are reported - never student names, phones or amounts -
so the report can be shared safely. Columns whose meaning is unclear (e.g. "2500", unnamed
columns that contain values, "old fees pending") are listed for a person to decide; they are
never mapped automatically.
"""
from __future__ import annotations

import re
import shutil
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .reader import _xls_to_xlsx, norm_key, norm_text

SYNONYMS = {
    "sr_no": ["sr. no.", "sr no", "sr.no", "sr. no", "srno", "serial no", "s. no.", "s.no"],
    "student_name": ["student name", "name of student", "student's name", "name"],
    "class_label": ["class", "std", "standard", "class/div", "class-div"],
    "contact": ["contact", "contact no", "contact no.", "mobile", "mobile no", "mobile no.", "phone", "parent mobile"],
    "total_fee": ["total amt", "total amount", "total fees", "total fee", "total"],
    "discount": ["disc. amt", "disc. amount", "disc amt", "discount", "concession", "disc."],
    "paid": ["paid amt", "paid amount", "paid", "amount paid"],
    "balance": ["balance", "balance amt", "balance amount", "bal.", "remaining"],
    "parent_name": ["parent name", "father name", "father's name", "guardian name", "parent's name"],
    "student_id": ["gr no", "gr no.", "g.r. no", "g.r. no.", "gr. no.", "admission no", "adm no", "adm. no."],
}
_LOOKUP = {norm_key(v): f for f, vs in SYNONYMS.items() for v in vs}


def _field(h) -> str | None:
    k = norm_key(h)
    return _LOOKUP.get(k) or _LOOKUP.get(k.rstrip("."))


def _header_row(rows: list[tuple]) -> int | None:
    best, best_n = None, 1
    for r, row in enumerate(rows[:8]):
        fields = {_field(h) for h in (row or ()) if h is not None} - {None}
        if "student_name" in fields and len(fields) > best_n:
            best, best_n = r, len(fields)
    return best


def inspect_folder(folder: str | Path, recursive: bool = True) -> dict:
    folder = Path(folder)
    files = sorted(p for p in (folder.rglob("*") if recursive else folder.iterdir())
                   if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm", ".xls") and not p.name.startswith(("~$", ".")))
    headings: dict[str, Counter] = defaultdict(Counter)  # field -> heading variants
    unknown: Counter = Counter()
    unnamed_with_values: Counter = Counter()
    header_rows: Counter = Counter()
    sheet_names, class_samples, problems = [], Counter(), []
    sheets_ok = 0
    for f in files:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / ("c" + f.suffix.lower())
            try:
                shutil.copy2(f, src)
                if f.suffix.lower() == ".xls":
                    conv = Path(tmp) / "c.xlsx"
                    _xls_to_xlsx(src, conv)
                    src = conv
                wb = load_workbook(src, read_only=True, data_only=True)
            except Exception as exc:
                problems.append(f"{f.name}: could not open ({exc.__class__.__name__})")
                continue
            try:
                for ws in wb.worksheets:
                    rows = list(ws.iter_rows(values_only=True))
                    h = _header_row(rows)
                    if h is None:
                        continue
                    sheets_ok += 1
                    header_rows[h + 1] += 1
                    sheet_names.append(ws.title)
                    head = rows[h]
                    data = [r for r in rows[h + 1:] if r and any(v not in (None, "") for v in r)]
                    for i, cell in enumerate(head):
                        filled = sum(1 for r in data if i < len(r) and r[i] not in (None, ""))
                        if cell is None or not norm_text(cell):
                            if filled:
                                unnamed_with_values[f"column {i + 1} (no heading)"] += filled
                            continue
                        fld = _field(cell)
                        if fld:
                            headings[fld][norm_text(cell)] += 1
                            if fld == "class_label":
                                for r in data[:50]:
                                    if i < len(r) and r[i]:
                                        class_samples[re.sub(r"\d", "9", norm_text(r[i]))] += 1
                        else:
                            unknown[norm_text(cell)] += 1
            finally:
                wb.close()
    return {"files": len(files), "class_sheets": sheets_ok, "heading_rows": dict(header_rows),
            "mapping": {k: [h for h, _ in v.most_common()] for k, v in headings.items()},
            "needs_decision": dict(unknown), "unnamed_columns_with_values": dict(unnamed_with_values),
            "sheet_name_examples": sorted(set(sheet_names))[:20], "class_value_shapes": dict(class_samples.most_common(10)),
            "problems": problems}


def suggest_toml(rep: dict) -> str:
    """A [layout]/[columns] block to paste into config.toml (mapping_confirmed stays false)."""
    hr = max(rep["heading_rows"], key=rep["heading_rows"].get) if rep["heading_rows"] else 2
    lines = ["[layout]", f"header_row = {hr}", f"first_data_row = {hr + 1}",
             "mapping_confirmed = false   # set true only after checking every line below", "test_mode = false", "",
             "[columns]"]
    for fld in SYNONYMS:
        hs = rep["mapping"].get(fld)
        if hs:
            val = f'"{hs[0]}"' if len(hs) == 1 else "[" + ", ".join(f'"{h}"' for h in hs) + "]"
            lines.append(f"{fld} = {val}")
    if rep["needs_decision"] or rep["unnamed_columns_with_values"]:
        lines.append("# NOT mapped - tell us what these mean before using them:")
        for h in list(rep["needs_decision"]) + list(rep["unnamed_columns_with_values"]):
            lines.append(f"#   {h}")
    return "\n".join(lines) + "\n"


def format_report(rep: dict) -> str:
    out = [f"Files: {rep['files']}   Class sheets found: {rep['class_sheets']}",
           f"Heading row(s): {rep['heading_rows']}", "", "Recognised columns:"]
    for fld, hs in rep["mapping"].items():
        out.append(f"  {fld:13} <- {' | '.join(hs)}")
    missing = [f for f in ("student_name", "balance", "contact") if f not in rep["mapping"]]
    if missing:
        out.append(f"  MISSING: {', '.join(missing)}")
    out += ["", "Columns that need YOUR decision (not used until you say what they mean):"]
    for h, n in {**rep["needs_decision"], **rep["unnamed_columns_with_values"]}.items():
        out.append(f"  - {h}  (in {n} {'sheets' if h in rep['needs_decision'] else 'rows'})")
    out += ["", f"Sheet names (examples): {', '.join(rep['sheet_name_examples'])}",
            f"Class values look like: {', '.join(rep['class_value_shapes'])}"]
    out += [f"Problem: {p}" for p in rep["problems"]]
    return "\n".join(out)
